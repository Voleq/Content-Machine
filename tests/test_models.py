import json

import pytest
from pydantic import ValidationError

from pipeline.models import (
    Candidate,
    CompanyData,
    CostReport,
    Lane,
    ShortScript,
    VisualPlanItem,
)


def test_short_script_valid(short_valid_json: str):
    script = ShortScript.model_validate(json.loads(short_valid_json))
    assert script.ticker == "EXMPL"
    assert len(script.headlines) == 2
    assert len(script.numbers) == 4
    assert all(len(r.values) == 6 for r in script.numbers), \
        "six periods: four fiscal years, the last full year, LTM"
    assert script.years == ["FY21", "FY22", "FY23", "FY24", "FY25", "LTM"]
    assert script.meme is not None and script.meme.key == "fomo-stages-wish-i-bought-doodle"
    assert "noise" in script.conclusion.lower()
    # The budget counts the SPOKEN text. Validated straight off the JSON like
    # this, `audio_script` still carries its inline tags — the parser is what
    # strips them, and it is where the budget is enforced — so measuring
    # `script.char_count` here counts `[PROP: …]` payloads as speech.
    from pipeline.models import SHORT_TAG_TYPES
    from pipeline.tagging import tokenize_tags

    clean, _tags, _warnings = tokenize_tags(script.audio_script,
                                            allowed=SHORT_TAG_TYPES)
    assert len(clean) <= 1200


def test_short_script_rejects_bad_row_index(fixtures_dir):
    raw = json.loads((fixtures_dir / "scripts" / "short_bad_index.json").read_text(encoding="utf-8"))
    with pytest.raises(ValidationError, match="row_index"):
        ShortScript.model_validate(raw)


def test_short_script_rejects_four_headlines(fixtures_dir):
    raw = json.loads((fixtures_dir / "scripts" / "short_bad_headlines.json").read_text(encoding="utf-8"))
    with pytest.raises(ValidationError, match="headlines"):
        ShortScript.model_validate(raw)


def test_short_script_requires_multi_year_numbers(short_valid_json: str):
    raw = json.loads(short_valid_json)
    raw["numbers"] = [{"label": "Revenue", "values": ["", "", "", "", "", "$496M"]}]  # single year
    with pytest.raises(ValidationError):
        ShortScript.model_validate(raw)


def test_short_script_conclusion_is_free_text(short_valid_json: str):
    """No verdict enum anywhere — the conclusion is prose."""
    raw = json.loads(short_valid_json)
    raw["conclusion"] = "Somewhere between noise and a cry for help."
    script = ShortScript.model_validate(raw)
    assert script.conclusion.startswith("Somewhere")
    import pipeline.models as models

    assert not hasattr(models, "Verdict"), "the verdict system is deleted"


def test_short_script_hook_length_cap(short_valid_json: str):
    raw = json.loads(short_valid_json)
    raw["hook_text"] = "x" * 91
    with pytest.raises(ValidationError):
        ShortScript.model_validate(raw)


def test_company_data_missing_classification():
    data = CompanyData(values={"company_name": "X", "price": 1.0})
    assert "ticker" in data.blocking_missing
    assert "pe_ttm" in data.warning_missing
    assert "pe_ttm" not in data.blocking_missing
    assert "short_interest" in data.warning_missing, "ownership is optional, never blocking"


def test_company_data_prompt_block_with_history():
    data = CompanyData(
        values={"company_name": "X", "price": 10.0},
        history_years=["FY-1", "FY-0"],
        history={"revenue": [100.0, 120.0], "fcf": [None, None]},
    )
    block = data.as_prompt_block()
    assert "[identity]" in block and "company_name = X" in block
    assert "[history" in block and "FY-1 | FY-0" in block
    assert "revenue: 100 | 120" in block
    assert "fcf" not in block.split("[history")[1], "all-empty history rows are omitted"
    assert data.has_history


def test_a_history_cell_keeps_every_digit():
    """`{:g}` kept six figures: revenue in raw dollars read 3.91035e+11."""
    data = CompanyData(values={"company_name": "X"}, history_years=["FY-1", "FY-0"],
                       history={"revenue": [391035000000.0, 1234.5678]})
    assert "revenue: 391035000000 | 1234.5678" in data.as_prompt_block()


def test_company_data_without_history():
    data = CompanyData(values={"company_name": "X"})
    assert not data.has_history
    assert "[history" not in data.as_prompt_block()


# --------------------------------------------------------------------------
# D6 - the template holds its yields as FRACTIONS and the fixture held
# percentages, so a thumbnail read the real workbook's FCF yield as "-0%".
# --------------------------------------------------------------------------


def test_every_percent_cell_in_the_template_is_a_fraction_here():
    """The shipped template is the data contract. A Snapshot cell it formats as
    a percent holds 0.074, not 7.4 - `D24` is `=D15/D22-1` under `0%` and
    `D43` divides the add-in's percent by 100 - and `D35` divides by 100 too,
    though its format is General."""
    import pathlib

    from openpyxl import load_workbook

    from pipeline.models import _FRACTION_FIELDS

    template = (pathlib.Path(__file__).resolve().parents[1] / "templates"
                / "dennis_data_template.xlsx")
    sheet = load_workbook(template)["Snapshot"]
    formatted = {sheet.cell(r, 2).value for r in range(1, sheet.max_row + 1)
                 if sheet.cell(r, 2).value and "%" in sheet.cell(r, 4).number_format}
    assert {"pct_from_52w_high", "debt_to_equity", "short_interest"} <= formatted
    assert formatted <= _FRACTION_FIELDS, formatted - _FRACTION_FIELDS
    assert "fcf_yield" in _FRACTION_FIELDS


def test_the_fixture_holds_fractions_like_the_template(fixtures_dir):
    from pipeline.company_data import load_company_data

    data = load_company_data(fixtures_dir / "company_data")
    assert data.values["fcf_yield"] == pytest.approx(-0.05)
    assert data.values["short_interest"] == pytest.approx(0.11)
    assert data.values["pct_from_52w_high"] == pytest.approx(-0.2983)
    # Debt / equity is a ratio shown as a percent: 1.4 is 140.0%.
    assert data.values["debt_to_equity"] == pytest.approx(1.4)
    assert data.dashboard["FCF yield"] == pytest.approx(-0.05)


def test_a_fraction_reaches_the_writer_as_a_percent():
    data = CompanyData(
        values={"fcf_yield": -0.05, "debt_to_equity": 1.4, "insider_own": 0.08,
                "ps_ttm": 62.0},
        dashboard={"FCF yield": -0.05, "% from 52w high": -0.2983,
                   "Net margin (LTM)": -18.0},
    )
    block = data.as_prompt_block()
    assert "fcf_yield = -5.0% (fraction -0.05)" in block
    assert "debt_to_equity = 140.0% (fraction 1.4)" in block
    assert "insider_own = 8.0% of shares outstanding (fraction 0.08)" in block
    assert "FCF yield: -5.0% (fraction -0.05)" in block
    assert "% from 52w high: -29.8% (fraction -0.2983)" in block
    # A multiple, and a margin that the History sheet holds in points, are
    # not fractions and are left exactly as they were.
    assert "ps_ttm = 62.0" in block
    assert "Net margin (LTM): -18.0" in block


def test_cost_report_render_text_short():
    report = CostReport(
        ticker="EXMPL",
        fmt="short",
        words=152,
        chars=784,
        tts_cached=False,
        est_tts_usd=0.12,
        headline_count=2,
        numbers_rows=4,
        numbers_years=5,
        est_render_minutes=0.7,
        mtd_spend_usd=3.10,
        monthly_cap_usd=50.0,
    )
    text = report.render_text()
    assert "EXMPL — SHORT — ready to render" in text
    assert "$0.12" in text and "$3.10 / $50.00" in text
    assert "Headlines: 2" in text and "4 rows × 5yr" in text
    assert report.approvable

    report.blocking.append("missing screenshot")
    assert not report.approvable
    assert "BLOCKED" in report.render_text()


def test_cost_report_render_text_long_visual_buckets():
    report = CostReport(
        ticker="EXMPL", fmt="long", words=1800, chars=9000,
        tts_cached=True, est_tts_usd=0.0,
        visuals=[
            VisualPlanItem(key="clown", kind="clip", source="library"),
            VisualPlanItem(key="q", kind="img", source="wikimedia"),
            VisualPlanItem(key="x", kind="clip", source="cache"),
            VisualPlanItem(key="y", kind="clip", source="filler"),
        ],
        filing_overlays=2, meme_count=1, meme_cap=2,
        mtd_spend_usd=0.0, monthly_cap_usd=50.0,
    )
    text = report.render_text()
    assert "Visuals: 4 (owned 1 / cache 1 / fetched 1 / filler 1)" in text
    assert "Filing overlays: 2" in text and "Memes: 1/2" in text


def test_candidate_why():
    c = Candidate(
        ticker="ABC",
        lane=Lane.TRENDING,
        score=0.9,
        reasons=["+18% today", "ST #3 trending", "vol 4× avg"],
    )
    assert c.why == "+18% today · ST #3 trending · vol 4× avg"


# --------------------------------------------------------------------------
# The approval report describes only what the short draws.
# --------------------------------------------------------------------------
# It named a chart style ("branded" or "hand-drawn napkin") and listed each
# scribble with its anchor. The short draws one price chart, the kit's, and no
# scribbles, and the approval screen is what the operator reads immediately
# before authorising the only spend in the pipeline.


def test_the_short_report_claims_no_chart_style_and_no_scribbles(settings):
    import json
    from pathlib import Path

    from pipeline.cost import SpendLedger, build_short_report
    from pipeline.parser_short import parse_short_script
    from pipeline.tts import TTSEngine

    raw = json.loads((Path(__file__).resolve().parents[1] / "fixtures" /
                      "scripts" / "short_valid.json").read_text(encoding="utf-8"))
    for style in ("clean", "marker"):
        raw["chart_style"] = style
        script, warnings = parse_short_script(json.dumps(raw), settings)
        assert script.annotations, "fixture: the script sets scribbles"
        text = build_short_report(script, warnings, settings,
                                  SpendLedger(settings),
                                  TTSEngine(settings)).render_text()
        assert not any(line.startswith(("Chart:", "Scribble"))
                       for line in text.splitlines())
        assert "napkin" not in text and "branded" not in text
