"""The diagnostic scripts an operator runs before a first live video (P5/P10).

These are not unit tests of library functions. They are the operator's
commands, and what a command is worth is its exit code — `scripts/check_sfx.py`
saying PASS while the audio gate would block a render is a worse state than
having no script at all, because it moves the discovery to the render.

So each one is run as a subprocess, both ways, against a throwaway tree.
Nothing here touches the network.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"


def _run(script: str, *args: str, **env_extra: str) -> subprocess.CompletedProcess:
    env = {"PATH": os.environ.get("PATH", ""), "PYTHONPATH": str(ROOT),
           "MOCK_MODE": "true"}
    env.update(env_extra)
    return subprocess.run([sys.executable, str(SCRIPTS / script), *args],
                          capture_output=True, text=True, env=env, cwd=ROOT,
                          timeout=180)


# --------------------------------------------------------------------------
# Every preflight script exists and behaves like one.
# --------------------------------------------------------------------------

# Grown by the commit that adds each script, so a name here always points at
# something that exists.
PREFLIGHT = ("check_sfx.py", "check_freshness.py", "check_llm_context.py",
             "backup_state.py", "check_preflight.py")


@pytest.mark.parametrize("name", PREFLIGHT)
def test_the_preflight_script_exists_and_is_runnable(name):
    """The README's preflight list names these. A step that points at a
    script that is not there is worse than a step that is missing."""
    path = SCRIPTS / name
    assert path.is_file(), f"{name} is named in the preflight and absent"
    text = path.read_text(encoding="utf-8")
    assert text.startswith("#!/usr/bin/env python3"), f"{name} has no shebang"
    assert '"""' in text.split("\n", 1)[1][:400], f"{name} has no docstring"
    # It must be importable-clean: a syntax error here is discovered by the
    # operator at the worst possible moment.
    import ast

    ast.parse(text)


# --------------------------------------------------------------------------
# check_freshness.py — the parser against the operator's real locale.
# --------------------------------------------------------------------------


def _workbook_with_as_of(dest: Path, value) -> Path:
    """The fixture workbook with its as-of cell rewritten."""
    from openpyxl import load_workbook

    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / "fixtures" / "company_data" / "dennis_data.xlsx", dest)
    wb = load_workbook(dest)
    ws = wb["Snapshot"]
    for row in ws.iter_rows():
        for cell in row:
            if str(cell.value).strip() == "as_of_date":
                ws.cell(row=cell.row, column=4, value=value)
    wb.save(dest)
    wb.close()
    return dest


def test_a_fresh_workbook_passes_the_freshness_check(tmp_path):
    """The passing case has to be reachable, or the script is just a wall."""
    import datetime as dt

    today = dt.date.today().isoformat()
    book = _workbook_with_as_of(tmp_path / "ws" / "dennis_data.xlsx", today)
    got = _run("check_freshness.py", str(book))
    assert got.returncode == 0, got.stdout + got.stderr
    assert "PASS" in got.stdout
    assert today in got.stdout


def test_a_stale_workbook_is_reported_as_stale_not_as_unreadable(tmp_path):
    """Two different problems with two different fixes, and the message the
    gate gives on its own does not distinguish them well enough to act on."""
    book = _workbook_with_as_of(tmp_path / "ws" / "dennis_data.xlsx",
                                "2020-01-01")
    got = _run("check_freshness.py", str(book))
    assert got.returncode == 1
    assert "2020-01-01" in got.stdout
    assert "READ fine" in got.stderr, "a stale date was reported as unparseable"
    assert "Refresh the workbook" in got.stderr


def test_a_date_the_parser_cannot_read_says_so_and_lists_what_it_accepts(
        tmp_path):
    """`DATA_STALE_BLOCKS` defaults to true and an unreadable date blocks
    too, so a locale the parser has never seen walls the operator off from
    their own pipeline with a message about staleness. This is the script's
    whole reason to exist: the sheet is whatever Capital IQ wrote."""
    book = _workbook_with_as_of(tmp_path / "ws" / "dennis_data.xlsx",
                                "10.IX.2026")
    got = _run("check_freshness.py", str(book))
    assert got.returncode == 1
    assert "NOTHING" in got.stdout
    assert "could not be PARSED" in got.stderr
    assert "%Y-%m-%d" in got.stderr, "it does not say what it would accept"
    assert "rather than turning the gate off" in got.stderr


@pytest.mark.parametrize("shape", [
    "2026-09-10",        # ISO, what the template asks for
    "09/10/2026",        # US-first
    "10/09/2026",        # day-first, a European locale
    "10-Sep-2026",
    "Sep 10, 2026",
    "10.09.2026",
    46276,               # a raw Excel serial
])
def test_the_shapes_a_capital_iq_export_might_write_all_read(tmp_path, shape):
    """Each of these is a real thing a locale writes into that cell, and the
    gate blocks on any it cannot read."""
    book = _workbook_with_as_of(tmp_path / f"ws{shape}" / "dennis_data.xlsx",
                                shape)
    got = _run("check_freshness.py", str(book))
    assert "NOTHING" not in got.stdout, (
        f"{shape!r} is unreadable to the gate:\n{got.stdout}{got.stderr}")


def test_the_script_explains_itself_with_no_arguments():
    got = _run("check_freshness.py")
    assert got.returncode == 2
    assert "check_freshness.py" in got.stderr


# --------------------------------------------------------------------------
# check_llm_context.py — the one that needs a live daemon.
# --------------------------------------------------------------------------


def test_the_context_check_reports_an_unreachable_daemon_as_unproven():
    """It needs a live Ollama, which the suite does not have. What matters
    offline is that it does not claim a PASS it cannot know: "nothing is
    proven" and "the marker came back" must never be confusable."""
    got = _run("check_llm_context.py", OLLAMA_BASE_URL="http://127.0.0.1:1")
    assert got.returncode == 2, got.stdout + got.stderr
    assert "Nothing is proven" in got.stdout or "Nothing is proven" in got.stderr
    assert "OK —" not in got.stdout


# --------------------------------------------------------------------------
# backup_state.py — the ledger is the only record of what was spent.
# --------------------------------------------------------------------------


def _seed_state(root: Path) -> Path:
    state = root / "state"
    (state / "jobs").mkdir(parents=True)
    (state / "spend.json").write_text(
        json.dumps({"2026-09": {"usd": 12.34, "tts_chars": 9000}}),
        encoding="utf-8")
    (state / "thesis.json").write_text('{"EXMPL": {}}', encoding="utf-8")
    (state / "jobs" / "abc123.json").write_text('{"id": "abc123"}',
                                                encoding="utf-8")
    return state


def test_the_backup_archives_state_and_can_be_restored(tmp_path):
    """P6: `state/` holds the spend ledger, and the ledger is the only record
    of what has been spent — so losing it loses the monthly cap silently.
    Asserted by unpacking the archive and comparing, not by trusting that
    tar was called."""
    import tarfile

    state = _seed_state(tmp_path)
    out = tmp_path / "backups"
    got = _run("backup_state.py", "--state", str(state), "--out", str(out))
    assert got.returncode == 0, got.stdout + got.stderr

    archives = list(out.glob("state-*.tar.gz"))
    assert len(archives) == 1, f"expected one archive, got {archives}"
    restored = tmp_path / "restored"
    restored.mkdir()
    with tarfile.open(archives[0]) as tar:
        tar.extractall(restored)

    back = restored / "state"
    assert json.loads((back / "spend.json").read_text())["2026-09"]["usd"] == 12.34
    assert (back / "thesis.json").is_file()
    assert (back / "jobs" / "abc123.json").is_file()
    # The ledger is called out by name, because it is the one whose loss
    # costs money rather than time.
    assert "ledger" in got.stdout and "2026-09" in got.stdout
    assert "Restore with:" in got.stdout


def test_the_backup_refuses_rather_than_writing_an_empty_archive(tmp_path):
    """An archive of nothing is worse than no archive: it is a backup you
    think you have."""
    empty = tmp_path / "state"
    empty.mkdir()
    got = _run("backup_state.py", "--state", str(empty),
               "--out", str(tmp_path / "b"))
    assert got.returncode == 1
    assert "empty" in got.stderr
    assert not list((tmp_path / "b").glob("*")) if (tmp_path / "b").is_dir() else True

    missing = _run("backup_state.py", "--state", str(tmp_path / "nope"),
                   "--out", str(tmp_path / "b2"))
    assert missing.returncode == 1
    assert "nothing to back up" in missing.stderr


def test_the_backup_lists_what_it_has_written(tmp_path):
    state = _seed_state(tmp_path)
    out = tmp_path / "backups"
    assert _run("backup_state.py", "--state", str(state),
                "--out", str(out)).returncode == 0
    listed = _run("backup_state.py", "--out", str(out), "--list")
    assert listed.returncode == 0
    assert "1 archive(s)" in listed.stdout
    assert "state-" in listed.stdout

    nothing = _run("backup_state.py", "--out", str(tmp_path / "elsewhere"),
                   "--list")
    assert nothing.returncode == 0
    assert "No archives" in nothing.stdout


# --------------------------------------------------------------------------
# check_preflight.py — the README's checklist, mirrored where a machine can
# answer it.
# --------------------------------------------------------------------------


def _fake_tree(root: Path, *, sfx_ok: bool, kit: bool) -> Path:
    """An assets tree in whatever state the test needs."""
    assets = root / "assets"
    (assets / "broll_library").mkdir(parents=True)
    sfx = assets / "sfx"
    sfx.mkdir()
    (sfx / "ding.wav").write_bytes(b"RIFF")
    if sfx_ok:
        from pipeline.audio_assets import AudioSource, save_sources

        save_sources(sfx, {"ding.wav": AudioSource(
            name="ding.wav", source="freesound.org/s/1/", licence="CC0",
            author="someone", generated=False)})
    if kit:
        shutil.copytree(ROOT / "assets" / "plates", assets / "plates")
    return assets


def test_the_preflight_blocks_on_the_sound_gate(tmp_path):
    """P0b is one of the two hard blockers, and the script has to say so
    rather than passing and leaving it for the render to discover."""
    assets = _fake_tree(tmp_path, sfx_ok=False, kit=False)
    got = _run("check_preflight.py", ASSETS_DIR=str(assets),
               STATE_DIR=str(tmp_path / "state"))
    assert got.returncode == 1
    assert "[FAIL] sound provenance" in got.stdout
    assert "fetch_sfx.py" in got.stdout
    assert "BLOCKED" in got.stderr


def test_the_preflight_blocks_on_a_missing_design_kit(tmp_path):
    """The other hard blocker. `assets/plates/` is a gitignored build
    product, and without it nothing renders on either lane."""
    assets = _fake_tree(tmp_path, sfx_ok=True, kit=False)
    got = _run("check_preflight.py", ASSETS_DIR=str(assets),
               STATE_DIR=str(tmp_path / "state"))
    assert got.returncode == 1
    assert "[FAIL] design kit" in got.stdout
    assert "ingest_kit.py" in got.stdout


@pytest.mark.skipif(not (ROOT / "assets" / "plates").is_dir(),
                    reason="the design kit is not built in this checkout")
def test_the_preflight_passes_once_both_blockers_are_cleared(tmp_path):
    """The passing case must be reachable, or the script is a wall rather
    than a check. Everything else it reports is advisory."""
    assets = _fake_tree(tmp_path, sfx_ok=True, kit=True)
    got = _run("check_preflight.py", ASSETS_DIR=str(assets),
               STATE_DIR=str(tmp_path / "state"))
    assert got.returncode == 0, got.stdout + got.stderr
    assert "PREFLIGHT OK" in got.stdout
    assert "[PASS] design kit" in got.stdout
    assert "[PASS] sound provenance" in got.stdout
    # …and it still says which steps remain a human's.
    assert "still yours" in got.stdout


def test_live_mode_checks_the_settings_that_only_matter_in_production(tmp_path):
    """`DELIVERY_BACKEND=local` writes a path and no link — correct for
    testing and silently useless in production — and an empty
    `SEC_USER_AGENT` costs three features quietly."""
    assets = _fake_tree(tmp_path, sfx_ok=True, kit=False)
    env = {"ASSETS_DIR": str(assets), "STATE_DIR": str(tmp_path / "state")}

    mocked = _run("check_preflight.py", **env)
    assert "[----] SEC_USER_AGENT" in mocked.stdout
    assert "[PASS] DELIVERY_BACKEND" in mocked.stdout

    live = _run("check_preflight.py", "--live", **env)
    assert "[FAIL] SEC_USER_AGENT" in live.stdout
    assert "[FAIL] DELIVERY_BACKEND" in live.stdout
    assert "silently useless in production" in live.stdout


def test_the_preflight_refuses_the_engine_with_no_production_mileage(tmp_path):
    """P4: `render_long_shots` is wired up and has never rendered a real
    video. A first live LONG must not be the one that finds out."""
    assets = _fake_tree(tmp_path, sfx_ok=True, kit=False)
    got = _run("check_preflight.py", ASSETS_DIR=str(assets),
               STATE_DIR=str(tmp_path / "state"),
               LONG_RENDER_ENGINE="shots")
    assert "[FAIL] LONG engine" in got.stdout
    assert "no production mileage" in got.stdout
    assert got.returncode == 1


def test_the_readme_preflight_and_the_script_name_the_same_scripts():
    """A checklist step pointing at a script that is not there is worse than
    a step that is missing, and the README is where the operator reads it."""
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    section = readme.split("## Preflight")[1].split("\n## ")[0]
    for name in ("check_preflight.py", "check_sfx.py", "check_freshness.py",
                 "check_llm_context.py", "ingest_kit.py", "fetch_sfx.py"):
        assert name in section, f"the preflight never mentions {name}"
        assert (SCRIPTS / name).is_file(), f"{name} is named and absent"
    # The two hard blockers are steps 1 and 3, and the sound gate is per file.
    assert "nothing renders on either lane" in section
    assert "room tone alone leaves fourteen" in section
    # And the deselect escape hatch, which does not clear the block.
    assert 'pytest -m "not audio_provenance"' in section
    assert "does not clear the render block" in section
